# -*- coding: utf-8 -*-
"""生成课程考核材料 4 个 xlsx：成员任务表 / 选题登记表 / 测试结果 / 作业04测试及风险分析"""
try:
    import openpyxl
except ImportError:
    import subprocess, sys
    subprocess.check_call([sys.executable, "-m", "pip", "install", "--quiet", "openpyxl>=3.1.0"])
    import openpyxl

from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from datetime import date

# ---- 配色（商务蓝） ----
XL_BORDER = "FFBFBFBF"
HEAD_FILL = PatternFill("solid", fgColor="FF4472C4")
HEAD_FONT = Font(bold=True, color="FFFFFFFF", size=11)
TITLE_FONT = Font(bold=True, size=14)
INPUT_FILL = PatternFill("solid", fgColor="FFD9E2F3")
SUM_FILL = PatternFill("solid", fgColor="FF2F5597")
SUM_FONT = Font(bold=True, color="FFFFFFFF")
THIN = Side(style="thin", color=XL_BORDER)
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top", horizontal="left")
CENTER = Alignment(horizontal="center", vertical="center")
GREEN_FILL = PatternFill("solid", fgColor="FFC6EFCE"); GREEN_FONT = Font(color="FF006100")
RED_FILL = PatternFill("solid", fgColor="FFFFC7CE"); RED_FONT = Font(color="FF9C0006")

OUT = r"D:\Ainterview\考核材料"

def style_header(ws, headers, row=1, title=None):
    r = row
    if title:
        ws.cell(row=row, column=1, value=title)
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(headers))
        ws.cell(row=row, column=1).font = TITLE_FONT
        ws.cell(row=row, column=1).alignment = CENTER
        r = row + 1
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=r, column=c, value=h)
        cell.fill = HEAD_FILL; cell.font = HEAD_FONT; cell.alignment = CENTER; cell.border = BORDER
    return r  # 表头行号

def fill_rows(ws, rows, start_row, widths, aligns=None):
    for i, row in enumerate(rows):
        r = start_row + i
        for c, v in enumerate(row, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.border = BORDER
            cell.alignment = WRAP if (aligns and aligns[c-1] == "wrap") else (CENTER if aligns and aligns[c-1] == "center" else Alignment(vertical="top"))
    for c, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(c)].width = w

# ================= 1. 成员任务表 =================
wb = openpyxl.Workbook(); wb.properties.title = "成员任务表"
ws = wb.active; ws.title = "成员任务"
headers = ["序号", "阶段", "任务名称", "负责人", "角色", "开始日期", "结束日期", "交付物", "状态", "备注"]
hr = style_header(ws, headers, title="AI智面（Ainterview）成员任务表")
data = [
    (1, "需求与设计", "需求分析与总体架构设计", "张嘉豪", "组长", date(2026,9,22), date(2026,9,26), "需求文档、架构图", "已完成", "确定2个感知模块+扩展能力选题"),
    (2, "需求与设计", "数据库与接口设计", "高弨", "开发", date(2026,9,24), date(2026,9,30), "ER图、API清单", "已完成", "evaluation表含voice/text evidence"),
    (3, "开发-迭代1", "共同业务闭环：登录/预约/存储/角色查询", "张嘉豪", "组长·开发", date(2026,9,28), date(2026,10,6), "业务闭环可运行", "已完成", "前后端与数据库真实联动"),
    (4, "开发-迭代1", "语音分析Python服务（librosa四维分析）", "高弨", "开发", date(2026,9,28), date(2026,10,5), "voice-analysis-service", "已完成", "语速/流利度/情绪稳定性/清晰度"),
    (5, "开发-迭代1", "语音分析Java接入与评分融合", "高弨", "开发", date(2026,10,4), date(2026,10,9), "VoiceAnalysisService等", "已完成", "按录音段聚合，旁路增强"),
    (6, "开发-迭代1", "简历解析与技能画像", "乔佳琪", "开发", date(2026,9,28), date(2026,10,6), "简历解析模块", "已完成", "扩展能力-简历与题目"),
    (7, "开发-迭代2", "智能组卷（按画像选题）", "乔佳琪", "开发", date(2026,10,6), date(2026,10,11), "组卷功能", "已完成", "现场可用新简历触发"),
    (8, "开发-迭代2", "文本理解与评价证据层（text_evidence）", "高弨", "开发", date(2026,10,10), date(2026,10,12), "V49迁移+评分证据", "已完成", "输入摘要/分档/规则版本可追溯"),
    (9, "开发-迭代2", "前端报告页：语音/文本逐题展示", "武心怡", "开发", date(2026,10,9), date(2026,10,12), "报告页区块", "已完成", "含未评估维度显示"),
    (10, "开发-迭代2", "面试大厅与作答交互页面", "武心怡", "开发", date(2026,10,2), date(2026,10,8), "前端页面", "已完成", "表单校验/加载/错误状态"),
    (11, "部署", "Docker Compose编排与健康检查", "张嘉豪", "组长·部署", date(2026,10,5), date(2026,10,12), "docker-compose.yml", "已完成", "6个服务一键启动"),
    (12, "部署", "环境配置模板与部署脚本", "乔佳琪", "部署", date(2026,10,8), date(2026,10,12), "config/.env.template", "已完成", "密钥占位符，不含真实密钥"),
    (13, "测试", "测试用例设计与执行", "武心怡", "测试", date(2026,10,9), date(2026,10,12), "测试结果.xlsx", "已完成", "正常/异常/权限用例"),
    (14, "测试", "项目风险分析与应对", "武心怡", "测试", date(2026,10,11), date(2026,10,12), "风险清单", "已完成", "5条风险含触发条件/影响/应对"),
    (15, "验收", "README与小组报告汇总", "张嘉豪", "组长", date(2026,10,12), date(2026,10,12), "README/小组报告", "已完成", "与现场版本一致"),
]
fill_rows(ws, data, hr + 1, [6, 12, 34, 10, 11, 12, 12, 22, 8, 30],
          ["center", "center", "wrap", "center", "center", "center", "center", "wrap", "center", "wrap"])
end = hr + len(data)
dv = DataValidation(type="list", formula1='"未开始,进行中,已完成,逾期"', allow_blank=True)
ws.add_data_validation(dv); dv.add(f"I{hr+1}:I{end+5}")
for r in range(hr+1, end+1):
    c = ws.cell(row=r, column=9)
    c.fill = GREEN_FILL if c.value == "已完成" else RED_FILL
ws.freeze_panes = f"A{hr+1}"
ws.auto_filter.ref = f"A{hr}:J{end}"

# ================= 2. 选题登记表 =================
wb2 = openpyxl.Workbook(); wb2.properties.title = "选题登记表"
ws2 = wb2.active; ws2.title = "选题登记"
headers2 = ["序号", "类别", "模块名称", "选择说明", "技术方案", "评分引擎接入方式", "报告追溯方式", "操作入口", "负责人"]
hr2 = style_header(ws2, headers2, title="AI智面（Ainterview）选题登记表（2个感知模块+1个扩展能力）")
data2 = [
    (1, "感知模块一", "语音分析", "对面试录音作答进行语音表现分析，满足'处理新输入并返回真实结果'的验收要求",
     "独立Python服务（FastAPI+librosa）：动态阈值语音掩码、语速、停顿比、能量/音高变异；webm经ffmpeg转码；Java按题目聚合录音段、多段按时长加权",
     "语音四维子分作为旁路增强维度写入evaluation表，参与融合报告",
     "evaluation.voice_evidence 存储四维原始特征与权重规则；报告页'逐题语音分析'逐题展示，未录音题目显示'未评估'",
     "候选人语音模式作答 → 报告页'逐题语音分析'", "高弨"),
    (2, "感知模块二", "文本理解与评价", "对文字回答做语义理解与分档评价，输出四维分与可追溯证据",
     "DeepSeek语义评分：分档规则(0-20未回答…93-100专家级)、四维评分与评语，校准后落库",
     "四维分数直接构成评分引擎主链路（专业/表达/逻辑/应变加权）",
     "evaluation.text_evidence 存储输入摘要与长度、命中分档及定义、四维原始分与校准分、评分规则版本号；报告页'逐题回答评价'展示'评价依据(可追溯)'",
     "候选人文字作答 → 报告页'逐题回答评价'", "高弨"),
    (3, "扩展能力", "简历与题目", "简历解析生成技能画像，按画像智能组卷，现场可用新简历触发新结果",
     "简历解析提取技能/项目画像；按画像从题库智能选题组卷并关联面试计划",
     "组卷题目进入面试与评分流程",
     "组卷记录关联面试与题目，报告题目来源可查",
     "候选人上传简历 → HR端'智能组卷'", "乔佳琪"),
]
fill_rows(ws2, data2, hr2 + 1, [6, 12, 14, 34, 40, 30, 44, 30, 10],
          ["center", "center", "center", "wrap", "wrap", "wrap", "wrap", "wrap", "center"])

# ================= 3. 测试结果.xlsx =================
wb3 = openpyxl.Workbook(); wb3.properties.title = "测试结果"
ws3 = wb3.active; ws3.title = "测试结果"
headers3 = ["序号", "测试类型", "功能模块", "测试输入", "预期结果", "实际结果", "是否通过", "执行人", "执行日期", "备注"]
hr3 = style_header(ws3, headers3, title="AI智面（Ainterview）测试结果表")
D = date(2026, 10, 12)
data3 = [
    (1, "正常", "登录与角色", "候选人/HR/管理员账号登录", "登录成功并跳转对应角色首页", "登录成功，角色路由正确", "通过", "武心怡", D, ""),
    (2, "正常", "预约面试", "候选人预约一场模拟面试", "预约成功，刷新后列表可见", "预约成功，数据库记录一致", "通过", "武心怡", D, ""),
    (3, "正常", "语音分析", "一段≥30s的webm录音作答", "报告逐题展示语音分与四维子分", "voice_score落库(29.5/39.8)，四维子分与证据完整", "通过", "高弨", D, "换新录音可复现"),
    (4, "正常", "文本理解与评价", "一段简答题文字回答", "报告逐题展示总分/分档/四维分/评语", "text_evidence落库(781字节)，含输入摘要与规则版本", "通过", "高弨", D, "换新文本可复现"),
    (5, "正常", "简历与题目", "上传一份PDF简历并触发智能组卷", "解析技能画像并生成题目", "解析成功，组卷题目与画像匹配", "通过", "乔佳琪", D, ""),
    (6, "正常", "融合报告", "完成一次含语音+文字的面试", "报告展示两个感知结果与融合评分，可追溯", "报告接口返回questionVoiceReviews/questionTextReviews", "通过", "武心怡", D, ""),
    (7, "正常", "报告发布查看", "候选人查看本人报告", "面试完成后可直接查看报告", "报告正常展示，含逐题语音/文本区块", "通过", "武心怡", D, "已去除发布限制"),
    (8, "正常", "通知发送", "管理员向候选人发送面试通知", "站内信送达，弹窗正常关闭", "站内信成功，邮件失败时降级不误报", "通过", "武心怡", D, "SMTP偶发SSL抖动已降级处理"),
    (9, "异常", "语音分析", "上传非音频文件(文本文件)作答", "明确报错提示，不产生评分", "返回'无法读取音频文件'错误提示", "通过", "高弨", D, "结果可复现"),
    (10, "异常", "评分任务", "DeepSeek接口超时", "任务重试，语音旁路失败不阻断", "重试生效，WARN日志，主链路正常", "通过", "高弨", D, ""),
    (11, "异常", "表单校验", "预约面试提交空表单/非法日期", "表单校验拦截并提示", "前端校验+后端参数校验双层拦截", "通过", "武心怡", D, ""),
    (12, "异常", "题库删除", "删除仍含题目的题库", "删除被拒绝并提示原因", "提示'题库下仍有题目时无法删除'", "通过", "武心怡", D, ""),
    (13, "权限", "越权访问报告", "候选人A访问候选人B的报告", "403拦截，结果可复现", "HTTP 403无权查看", "通过", "武心怡", D, ""),
    (14, "权限", "企业端越权", "企业账号访问非授权候选人报告", "拒绝访问", "权限拦截生效", "通过", "乔佳琪", D, ""),
    (15, "权限", "管理员接口", "候选人调用管理员重新评分接口", "403拒绝", "PreAuthorize拦截生效", "通过", "武心怡", D, ""),
]
fill_rows(ws3, data3, hr3 + 1, [6, 9, 15, 30, 30, 36, 9, 9, 12, 24],
          ["center", "center", "center", "wrap", "wrap", "wrap", "center", "center", "center", "wrap"])
end3 = hr3 + len(data3)
for r in range(hr3+1, end3+1):
    c = ws3.cell(row=r, column=7)
    c.fill = GREEN_FILL if c.value == "通过" else RED_FILL
ws3.freeze_panes = f"A{hr3+1}"
ws3.auto_filter.ref = f"A{hr3}:J{end3}"

# ================= 4. 作业04 测试及项目风险分析 =================
wb4 = openpyxl.Workbook(); wb4.properties.title = "作业04 测试及项目风险分析"
ws4 = wb4.active; ws4.title = "测试用例"
headers4 = ["序号", "测试类型", "测试项", "测试输入/步骤", "预期结果", "实际结果", "执行证据"]
hr4 = style_header(ws4, headers4, title="作业04 —— 测试用例（3正常+2异常+1权限）")
data4 = [
    (1, "正常", "候选人预约面试", "登录候选人账号→选择题库与时间→提交预约→刷新查询", "预约成功且列表可见，数据库记录一致", "通过", "运行截图：预约列表"),
    (2, "正常", "语音分析感知模块", "以语音模式回答一题(≥30s录音)→触发评分→查看报告", "报告逐题展示语音分与语速/流利度/情绪稳定性/清晰度四维子分", "通过", "运行截图+evaluation表记录"),
    (3, "正常", "文本理解与评价感知模块", "以文字回答一道简答题→触发评分→查看报告", "报告逐题展示总分/分档/四维分/评语及评价依据", "通过", "运行截图+text_evidence记录"),
    (4, "异常", "非音频文件上传", "语音作答时上传文本文件", "明确报错提示，不产生评分结果", "通过", "错误提示截图"),
    (5, "异常", "DeepSeek服务超时", "模拟外部模型超时后触发评分", "任务重试，语音旁路失败不阻断评分主链路", "通过", "backend WARN日志"),
    (6, "权限", "候选人越权访问他人报告", "候选人A用API直接请求候选人B的报告接口", "HTTP 403拦截，结果可复现", "通过", "403响应截图+复现步骤"),
]
fill_rows(ws4, data4, hr4 + 1, [6, 8, 20, 36, 34, 8, 22],
          ["center", "center", "center", "wrap", "wrap", "center", "wrap"])

ws4b = wb4.create_sheet("项目风险")
headers4b = ["序号", "风险描述", "触发条件", "影响", "应对措施", "责任人", "状态"]
hr4b = style_header(ws4b, headers4b, title="作业04 —— 项目风险分析（5条）")
data4b = [
    (1, "外部大模型服务不可用导致评分失败", "DeepSeek接口超时、限流或密钥失效", "评分延迟或任务失败", "任务重试机制+超时降级；管理员可重新触发评分", "高弨", "已落地"),
    (2, "长音频导致语音分析慢或崩溃", "录音超过3分钟或并发分析多个大文件", "语音维度缺失或响应超时", "pyin仅取前30秒+readTimeout提升至60s+旁路降级为'未评估'", "高弨", "已落地"),
    (3, "部署环境差异导致服务启动失败", "端口占用、卷挂载路径不一致、镜像拉取失败", "现场演示无法启动", "docker compose统一编排+健康检查+本地镜像兜底", "张嘉豪", "已落地"),
    (4, "录音格式兼容性问题", "浏览器MediaRecorder产出webm/opus", "librosa无法直读，语音分析失败", "ffmpeg转码兜底（直读失败自动转wav重试）", "高弨", "已落地"),
    (5, "现场演示数据缺失", "验收现场无有效音频/简历测试材料", "感知模块无法演示真实结果", "data/测试数据目录自备样例音频与简历，现场可换新输入复现", "武心怡", "已就绪"),
]
fill_rows(ws4b, data4b, hr4b + 1, [6, 30, 30, 24, 40, 10, 10],
          ["center", "wrap", "wrap", "wrap", "wrap", "center", "center"])

wb.save(OUT + r"\小组材料\成员任务表.xlsx")
wb2.save(OUT + r"\小组材料\选题登记表.xlsx")
wb3.save(OUT + r"\小组材料\测试结果.xlsx")
wb4.save(OUT + r"\个人材料\作业04-学号-武心怡-测试及项目风险分析.xlsx")
print("OK: 4 workbooks generated")
